"""Инвентаризационен опис в Excel: таблицата „под 20 лв“ и „над 20 лв“ с колона за всеки картон.

Форматът следва работната таблица на потребителя („РАБОТЕН ТАБЛИЦА EXCEL И АЖУР ИНВЕНТАР КАРТОНИ …“):
ред 9 е заглавието на колоните (имената са завъртени вертикално), ред 10 – номерата на картоните,
от ред 11 са позициите от склада в реда на импорта. Колоните след картоните са с формули:
ПО КАРТОН, В СКЛАД, В Т.Ч. ЗА БРАК, ФАКТИЧЕСКА НАЛИЧНОСТ, ИЗЛИШЪЦИ, ЛИПСИ.

Зачисленото по картон е по инв. номер; ако номерът е на няколко позиции в склада, се разпределя като в
„Наличности“ (първо под прага, позицията „авариен резерв“ последна), а остатъкът отива на последната позиция.
"""
import io
import re
from collections import defaultdict
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from db import get_db

FONT = "Sofia Sans"
# Ред на колоните, както е в таблицата на потребителя; служителите извън списъка са след тях по азбучен ред.
COLUMN_ORDER = [
    "Онник Гюлян", "Ангел Милев", "Петър Коилов", "Николай Папалаков", "Стоян Коджаманов", "Величка Пелтекова",
    "Павел Павлов", "Спаска Сокерова", "Иван Марков", "Петрунка Николова", "Силвия Чучулигова", "Ивелин Иванов",
    "Стоян Иванов", "Красимир Николов", "Данаил Щерев", "Анна Челебиева", "Димитър Костадинов", "Димо Седянков",
    "Ангел Чучулигов", "Наско Тарашев", "Костадин Павлов", "Исмаил Татаров", "Борис Коилов", "Ангел Танин",
    "Петър Серафимов", "Юсуф Сюлейманов", "Север Мюмюнов",
]
TOTALS = ["ПО КАРТОН", "В СКЛАД", "В Т.Ч. ЗА БРАК", "ФАКТИЧЕСКА НАЛИЧНОСТ", "ИЗЛИШЪЦИ", "ЛИПСИ"]
TOTAL_WIDTHS = [5.86, 8.5, 5.86, 10.14, 5.86, 6.71]
# Групите на номерата без 9- (1000 А, 2001 Б …): буквата е в първата колона на първия ред от групата.
LETTERS = "АБВГДЕЖЗИЙКЛМНОПРСТУФХЦЧШЩЪЮЯ"
GROUP_LABELS = {33: 'юбилей язовирна стена "Въча"'}
SHEETS = [("under", "Под 20лв"), ("over", "Над 20лв")]

THIN = Side(style="thin")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _font(bold=False, size=10):
    return Font(name=FONT, size=size, bold=bold)


def _holders(db, warehouse_id):
    """Колоните с картони: личните (по COLUMN_ORDER), после по съоръжения и автомобилите."""
    from cards import EFF
    rows = db.execute(
        "SELECT c.id, c.kind, c.title, e.first_name, e.last_name FROM cards c LEFT JOIN employees e ON e.id = c.employee_id"
        f" WHERE EXISTS (SELECT 1 FROM card_items ci WHERE ci.card_id = c.id AND {EFF} = ? AND ci.section = 'inventory'"
        "  AND ci.returned_on IS NULL)"
        "  OR (c.kind = 'personal' AND e.left_on IS NULL)", (warehouse_id,)).fetchall()
    holders = []
    for r in rows:
        label = f"{r['first_name']} {r['last_name']}" if r["kind"] == "personal" else r["title"]
        rank = COLUMN_ORDER.index(label) if label in COLUMN_ORDER else len(COLUMN_ORDER)
        holders.append({"card_id": r["id"], "label": label,
                        "key": ({"personal": 0, "facility": 1, "vehicle": 2}[r["kind"]], rank, label.casefold())})
    return sorted(holders, key=lambda h: h["key"])


def _issued(db, warehouse_id):
    """(картон, инв. номер) -> зачислено."""
    from cards import EFF
    result = defaultdict(float)
    for r in db.execute(
            "SELECT ci.card_id, ci.inv_number, SUM(ci.quantity) AS q FROM card_items ci JOIN cards c ON c.id = ci.card_id"
            f" WHERE {EFF} = ? AND ci.section = 'inventory' AND ci.returned_on IS NULL AND ci.inv_number != ''"
            " GROUP BY ci.card_id, ci.inv_number", (warehouse_id,)):
        result[(r["card_id"], r["inv_number"])] += r["q"]
    return result


def _reserve_need(warehouse_id):
    """Инв. номер -> нужното за аварийния резерв (в мярката на склада)."""
    from reserve import coverage
    need = defaultdict(float)
    for item in coverage(warehouse_id):
        for n in item["linked"]:
            need[n] += item["quantity"] * item["factor"] / len(item["linked"])
    return need


def _spread(rows, amounts):
    """Разпределя количествата [(колона, количество)] на един инв. номер между позициите му.

    Редът е като в cards.allocate: първо под прага, позицията „резерв“ последна; остатъкът е на последната.
    Връща {(id на позиция, колона): количество}.
    """
    rows = sorted(rows, key=lambda r: ("РЕЗЕРВ" in r["name"].upper(), r["value_class"] != "under", r["id"]))
    room = {r["id"]: r["quantity"] for r in rows}
    result = {}
    for column, left in amounts:
        for i, row in enumerate(rows):
            take = left if i == len(rows) - 1 else min(left, max(room[row["id"]], 0))
            if take:
                result[(row["id"], column)] = result.get((row["id"], column), 0) + take
                room[row["id"]] -= take
                left -= take
            if not left:
                break
    return result


def _order(row):
    """Първо номерата 9-… (по номер), после азбучните групи (1000 А, 2001 Б …) и следващите по код."""
    inv = row["inv_number"]
    stocked = inv.startswith(("9-", "09-"))
    digits = tuple(int(d) for d in re.findall(r"\d+", inv.split("-", 1)[1] if stocked else inv))
    return (not stocked, digits, inv, row["id"])


def _qty(value):
    return int(value) if float(value).is_integer() else round(value, 3)


def _label(row, mol):
    """„1;ПАВЕЛ ПАВЛОВ;9-020003;БРАДВИ“ или „1;ПАВЕЛ ПАВЛОВ;2001 БРАДВА“, както в импорта."""
    head = f"{row['warehouse_id']};{mol}"
    if row["inv_number"].startswith(("9-", "09-")):
        return f"{head};{row['inv_number']};{row['name']}"
    return f"{head};{row['inv_number']} {row['name']}".rstrip()


def build(warehouse_id, as_of=None):
    """Връща (име на файла, байтове) на описа за склада."""
    from stock import warehouses
    db = get_db()
    as_of = as_of or date.today()
    holders = _holders(db, warehouse_id)
    issued = _issued(db, warehouse_id)
    need = _reserve_need(warehouse_id)
    stock = db.execute(
        "SELECT s.*, e.first_name, e.last_name FROM stock_items s LEFT JOIN employees e ON e.id = s.mol_employee_id"
        " WHERE s.warehouse_id = ? ORDER BY s.id", (warehouse_id,)).fetchall()
    stock.sort(key=_order)
    by_inv = defaultdict(list)
    for row in stock:
        by_inv[row["inv_number"]].append(row)
    cells = {}
    for inv, rows in by_inv.items():
        amounts = [(h["card_id"], issued[(h["card_id"], inv)]) for h in holders if issued.get((h["card_id"], inv))]
        if need.get(inv):
            amounts.append(("reserve", need[inv]))
        cells.update(_spread(rows, amounts))

    mol = next((f"{r['first_name']} {r['last_name']}" for r in stock if r["first_name"]), "")
    wh_name = next((w["name"] for w in warehouses() if w["id"] == warehouse_id), "")
    wb = Workbook()
    wb.remove(wb.active)
    for value_class, title in SHEETS:
        _sheet(wb.create_sheet(title), [r for r in stock if r["value_class"] == value_class], holders, cells,
               value_class == "over", mol, warehouse_id, as_of)
    out = io.BytesIO()
    wb.save(out)
    return f"Инвентаризационен опис {wh_name} {as_of.strftime('%d.%m.%Y')}.xlsx", out.getvalue()


def _sheet(ws, rows, holders, cells, with_price, mol, warehouse_id, as_of):
    columns = ["№ по ред", "НАИМЕНОВАНИЕ", "Мярка", "СЧЕТОВОДНА НАЛИЧНОСТ"] + (["ЕДИНИЧНА ЦЕНА"] if with_price else [])
    first_holder = len(columns) + 1
    holder_cols = [(h["card_id"], h["label"]) for h in holders] + [("reserve", "АВАРИЕН РЕЗЕРВ")]
    columns += [label for _, label in holder_cols]
    last_card = first_holder + len(holders) - 1  # ПО КАРТОН е до аварийния резерв, без него
    reserve = get_column_letter(last_card + 1)
    total_col = len(columns) + 1
    columns += TOTALS
    col = {name: get_column_letter(total_col + i) for i, name in enumerate(TOTALS)}

    ws["B1"], ws["B3"] = 'ЯЗОВИРЕН  РАЙОН  "ВЪЧА"', "ИНВЕНТАРИЗАЦИОНЕН  ОПИС"
    ws["B5"] = f"на материални запаси към {as_of.strftime('%d.%m.%Y')} г."
    ws["B7"] = f"Сметка 98/01 Склад {warehouse_id} {mol}"
    for addr in ("B1", "B3"):
        ws[addr].font = _font(True, 11)
    for addr in ("B5", "B7"):
        ws[addr].font = _font(True)

    for i, name in enumerate(columns, start=1):
        cell = ws.cell(9, i, name)
        cell.font, cell.border = _font(True), BOX
        vertical = i >= first_holder and name != "ЕДИНИЧНА ЦЕНА"
        cell.alignment = Alignment(horizontal="center", vertical="bottom" if vertical else "center",
                                   text_rotation=90 if vertical else 0, wrap_text=True)
        number = ws.cell(10, i, i - first_holder + 1 if first_holder <= i <= last_card else None)
        number.font, number.border, number.alignment = _font(True), BOX, Alignment(horizontal="center")
    ws.row_dimensions[9].height = 138.75
    ws.freeze_panes = "E11" if not with_price else "F11"

    r, n, group = 11, 0, None
    for row in rows:
        stocked = row["inv_number"].startswith(("9-", "09-"))
        first_cell = None
        if stocked:
            n += 1
            first_cell = n
        else:
            if group is None:
                r += 2  # празни редове преди номерата по азбучните групи
            code = int(re.match(r"\d+", row["inv_number"]).group()) if re.match(r"\d+", row["inv_number"]) else 0
            new_group = code // 1000 if code < 100000 else code // 100000  # „400001“ е 4-та група
            if new_group != group:
                group = new_group
                first_cell = GROUP_LABELS.get(group) or (LETTERS[group - 1] if 0 < group <= len(LETTERS) else None)
        values = [first_cell, _label(row, f"{row['first_name']} {row['last_name']}".upper() if row["first_name"] else ""),
                  row["unit"], _qty(row["quantity"]) if row["quantity"] or row["name"] else None] + ([None] if with_price else [])
        values += [_qty(cells[(row["id"], key)]) if cells.get((row["id"], key)) else None for key, _ in holder_cols]
        for i, value in enumerate(values, start=1):
            cell = ws.cell(r, i, value)
            cell.font, cell.border = _font(bold=(i == 1 and not stocked)), BOX
            if i != 2:
                cell.alignment = Alignment(horizontal="center")
        ws.cell(r, 2).number_format = "@"
        ws.cell(r, 3).number_format = "@"
        if stocked:
            cards_sum = f"SUM({get_column_letter(first_holder)}{r}:{get_column_letter(last_card)}{r})" if holders else "0"
            # В СКЛАД = счетоводна наличност − авариен резерв − по картон; фактическата ги събира обратно
            formulas = [f"={cards_sum}", f"=D{r}-{reserve}{r}-{col['ПО КАРТОН']}{r}", None,
                        f"={col['ПО КАРТОН']}{r}+{reserve}{r}+{col['В СКЛАД']}{r}", f"={col['ФАКТИЧЕСКА НАЛИЧНОСТ']}{r}-D{r}",
                        f"=D{r}-{col['ФАКТИЧЕСКА НАЛИЧНОСТ']}{r}"]
        else:
            formulas = [None] * len(TOTALS)
        for i, value in enumerate(formulas):
            cell = ws.cell(r, total_col + i, value)
            cell.font, cell.border, cell.alignment = _font(), BOX, Alignment(horizontal="center")
        r += 1

    widths = {"A": 5.29, "B": 86.57 if with_price else 63.29, "C": 6.71, "D": 13.71}
    if with_price:
        widths["E"] = 9
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width
    for i in range(first_holder, last_card + 1):
        ws.column_dimensions[get_column_letter(i)].width = 3.43
    ws.column_dimensions[reserve].width = TOTAL_WIDTHS[0]
    for i, width in enumerate(TOTAL_WIDTHS):
        ws.column_dimensions[get_column_letter(total_col + i)].width = width
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "9:10"
